# sheet_manager.py
import os
from typing import Dict, Any, Optional
from openpyxl import Workbook, load_workbook


class SpreadsheetManager:
    """
    Manages reading and writing data to an Excel spreadsheet.

    This class is responsible for:
    - Creating a new Excel file if it does not exist
    - Loading an existing Excel file
    - Appending structured OCR data into the spreadsheet
    - Saving changes to disk
    """

    def __init__(self, excel_file: str) -> None:
        """
        Initialize the spreadsheet manager.

        Args:
            excel_file (str): Path to the Excel file.
        """
        self.excel_file: str = excel_file

        if os.path.exists(self.excel_file):
            self.wb = load_workbook(self.excel_file)
            self.ws = self.wb.active
        else:
            self.wb = Workbook()
            self.ws = self.wb.active

            # Header row (only created once)
            self.ws.append([
                "Data",
                "Hora",
                "S",
                "W",
                "Km",
                "Lado",
                "Foto"
            ])

            self.save_file()

    def update(self, data: Optional[Dict[str, Any]]) -> None:
        """
        Append a new row of extracted OCR data into the spreadsheet.

        Args:
            data (Optional[Dict[str, Any]]): Dictionary containing OCR-extracted fields.
                Expected keys:
                - date (str)
                - time (str)
                - s (str)
                - w (str)
                - km (str)
                - side (str)
                - new_name (str)
        """
        if not data:
            print("Error: data not found!")
            return

        self.ws.append([
            data.get("date", ""),
            data.get("time", ""),
            data.get("s", ""),
            data.get("w", ""),
            data.get("km", ""),
            data.get("side", ""),
            data.get("new_name", "")
        ])

        print(f"OK -> {data.get('new_name', '')}")

        self.save_file()

    def save_file(self) -> None:
        """
        Save the current workbook to disk.
        """
        self.wb.save(self.excel_file)
        print("\nProcessing complete!")
